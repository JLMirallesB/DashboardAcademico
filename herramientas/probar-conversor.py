"""Las cifras del conversor, recalculadas por un Excel de verdad.

    python3 herramientas/probar-conversor.py

Necesita Microsoft Excel instalado, y por eso no va en `pruebas/`: la red de
node no puede evaluar un LET. Lo que hace:

  1. inventa una exportación con la forma de la real —ni un dato de nadie—
     y con cada caso raro que se ha visto o se puede ver: extraordinarias
     vacías, un 1, una pendiente de otro curso, una doble especialidad, una
     asignatura que no está en el catálogo, un curso sin etapa, una nota con
     letras, una evaluación que no toca;
  2. genera un conversor con esas filas en ENTRADA, en una carpeta temporal;
  3. lo abre en Excel, recalcula, guarda y cierra;
  4. calcula aquí, a mano y por otro camino, lo que tendría que salir, y lo
     compara celda a celda.

Si el paso 4 y Excel no coinciden, el que está mal es el libro — o esta
prueba, y entonces hay que saber cuál antes de tocar nada.
"""
import json, os, random, subprocess, sys, tempfile
sys.path.insert(0, os.path.dirname(__file__))
import importlib.util
spec = importlib.util.spec_from_file_location('gen', os.path.join(os.path.dirname(__file__), 'generar-conversor.py'))
gen = importlib.util.module_from_spec(spec); spec.loader.exec_module(gen)
import openpyxl

random.seed(7)
CAB = gen.CABECERAS_ENTRADA

def fila(nia, curso_mat, curso_cont, ev, asig, nota, tipo='Normal', codigo='100'):
    d = dict.fromkeys(CAB, '')
    d.update(anoacademico_id='2025', codcentro='12000000', denominacionespecifica='CENTRO DE PRUEBA',
             codensenanza='10', curso_mat=curso_mat, curso_cont=curso_cont,
             nombre='NOMBRE%s' % nia, apellido1='APELLIDO%s' % nia, apellido2='', nia=nia, sexo='H',
             ev_codigo=ev, ev_abrv=ev, ev_desc='EVALUACIÓN', ev_orden='1', ev_obligatoria='1',
             ev_para_fct_proyecto='0', co_codigo=codigo, co_desc=asig, co_estutoria='0',
             co_esproyecto='0', co_caracter='Común', co_tipobasico=tipo, co_singularidad='',
             nota='' if nota is None else nota,
             aprueba='' if nota is None or not str(nota).isdigit() else ('1' if int(nota) >= 5 else '0'),
             promociona='1')
    return [d[c] for c in CAB]

filas = []
nia = 1000
# Elemental: FI, con Lenguaje Musical, Coro, su instrumento y Conjunto en 3.º y 4.º
for curso in ['1EEM', '2EEM', '3EEM', '4EEM']:
    for _ in range(4):
        nia += 1
        instr = random.choice(['PIANO', 'VIOLÍN', 'CLARINETE', 'PERCUSIÓN', 'DULZAINA'])
        for a in ['LENGUAJE MUSICAL', 'CORO', instr] + (['CONJUNTO'] if curso in ('3EEM', '4EEM') else []):
            filas.append(fila(nia, curso, curso, 'FI', a, random.choice([1, 4, 5, 6, 7, 8, 9, 10])))
# Profesional: OR y EX de cada asignatura; la EX solo trae nota si la OR suspendió
COMUNES = {'1EPM': ['LENGUAJE MUSICAL', 'ORQUESTA/BANDA/CONJUNTO', 'CONJUNTO'],
           '3EPM': ['ARMONÍA', 'MÚSICA DE CÁMARA', 'PIANO COMPLEMENTARIO'],
           '5EPM': ['ANÁLISIS', 'HISTORIA DE LA MÚSICA', 'COMPLEMENTO PIANÍSTICO / CLAVECINÍSTICO']}
def epm(nia, curso, instrumentos, extra=()):
    for a in list(instrumentos) + COMUNES[curso] * len(instrumentos) + list(extra):
        cont, tipo = curso, 'Normal'
        if isinstance(a, tuple):
            a, cont, tipo = a
        o = random.choice([1, 3, 5, 6, 7, 8, 9, 10])
        filas.append(fila(nia, curso, cont, 'OR', a, o, tipo))
        filas.append(fila(nia, curso, cont, 'EX', a, random.choice([4, 5, 6]) if o < 5 else None, tipo))
for curso in COMUNES:
    for _ in range(4):
        nia += 1
        epm(nia, curso, [random.choice(['PIANO', 'TROMPETA', 'GUITARRA ELÉCTRICA', 'CANTO'])])
# Casos raros, uno de cada
nia += 1; epm(nia, '5EPM', ['VIOLÍN'], extra=[('ARMONÍA', '4EPM', 'Pendiente')])   # pendiente
nia += 1; epm(nia, '3EPM', ['PIANO', 'VIOLONCELLO'])                               # doble especialidad
nia += 1; epm(nia, '1EPM', ['BANDURRIA'])                     # especialidad que no está en el catálogo
filas.append(fila(nia, '1', '1', 'OR', 'PIANO', 7))                                # curso sin etapa
filas.append(fila(nia + 1, '2EEM', '2EEM', 'FI', 'CORO', 'NP'))                   # nota con letras
filas.append(fila(nia + 1, '2EEM', '2EEM', '3EV', 'CORO', 6))                     # evaluación que no toca
filas.append(fila(nia + 1, '2EEM', '2EEM', 'FI', 'Lenguaje  Musical ', 8))        # espacios de más

# ── Lo esperado, calculado aquí por otro camino ────────────────────────────

cat = gen.catalogo()
canon = {(a.lower(), e): a for a, e, _ in cat}
esp = {(a.lower(), e) for a, e, s in cat if s == 'Sí'}
col = {c: i for i, c in enumerate(CAB)}
def v(r, c): return str(r[col[c]]).strip()
def etapa(r): return v(r, 'curso_cont')[-3:]
def nombre(r): return ' '.join(v(r, 'co_desc').split()) if False else v(r, 'co_desc')

# TRIM de Excel colapsa los espacios interiores; strip de Python no.
def trim(s): return ' '.join(str(s).split())

ids = {}
for r in filas:
    ids.setdefault(v(r, 'nia'), len(ids) + 1)
primera_esp = {}
for r in filas:
    k = (trim(v(r, 'co_desc')).lower(), etapa(r))
    if k in esp:
        primera_esp.setdefault((v(r, 'nia'), etapa(r)), canon[k])

def esperado(et):
    salida = []
    for r in filas:
        t = v(r, 'nota')
        if etapa(r) != et or t == '' or not t.replace(',', '').replace('.', '').isdigit():
            continue
        k = (trim(v(r, 'co_desc')).lower(), et)
        nombre = canon.get(k, trim(v(r, 'co_desc')))
        instr = nombre if k in esp else primera_esp.get((v(r, 'nia'), et), '')
        n = float(t.replace(',', '.'))
        apto = 'S' if v(r, 'aprueba') == '1' else 'N'
        salida.append([ids[v(r, 'nia')], None, None, v(r, 'curso_cont'), instr or None, None,
                       v(r, 'ev_codigo'), v(r, 'co_codigo'), nombre, n, n, apto, None, None])
    return salida

# ── Excel ──────────────────────────────────────────────────────────────────

carpeta = tempfile.mkdtemp(prefix='conversor-')
ruta = os.path.join(carpeta, 'conversor-con-datos.xlsx')
datos_json = os.path.join(carpeta, 'filas.json')
json.dump(filas, open(datos_json, 'w'))
gen.construir(ruta, filas)
print('Generado con %d filas: %s' % (len(filas), ruta))

script = '''
tell application "Microsoft Excel"
    set wb to open workbook workbook file name (POSIX file "%s" as text)
    calculate full
    save wb
    close wb saving no
end tell''' % ruta
r = subprocess.run(['osascript', '-e', script], capture_output=True, text=True, timeout=300)
if r.returncode != 0:
    print('Excel no ha podido abrirlo:', r.stderr.strip()); sys.exit(2)

libro = openpyxl.load_workbook(ruta, data_only=True)
fallos = 0
def comprobar(titulo, ok, detalle=''):
    global fallos
    print(('  ✓ ' if ok else '  ✗ ') + titulo + ('' if ok else '  → ' + str(detalle)))
    fallos += 0 if ok else 1

def valores(hoja):
    h = libro[hoja]
    out = []
    for fila_ in h.iter_rows(min_row=2, max_col=14, values_only=True):
        if all(x in (None, '') for x in fila_):
            break
        out.append([None if x == '' else x for x in fila_])
    return out

for et in ('EEM', 'EPM'):
    real, esp_ = valores('DATOS_' + et), esperado(et)
    comprobar('DATOS_%s: %d filas' % (et, len(esp_)), len(real) == len(esp_), '%d en Excel' % len(real))
    distintas = [(i + 2, a, b) for i, (a, b) in enumerate(zip(real, esp_))
                 if [str(x) if x is not None else None for x in a] != [str(x) if x is not None else None for x in b]
                 and not all((x == y) or (isinstance(x, (int, float)) and isinstance(y, (int, float)) and abs(x - y) < 1e-9)
                             or (x is None and y is None) for x, y in zip(a, b))]
    comprobar('DATOS_%s: celda a celda' % et, not distintas, distintas[:2])

control = {libro['CONTROL'].cell(i, 1).value: libro['CONTROL'].cell(i, 2).value
           for i in range(3, libro['CONTROL'].max_row + 1)}
print('\nCONTROL:')
for k, x in control.items():
    print('   %-55s %s' % (k, x))
comprobar('no falta ninguna columna', control['Columnas que faltan en ENTRADA'] == 'Ninguna')
comprobar('filas sin nota', control['Filas sin nota, descartadas'] == sum(1 for r in filas if v(r, 'nota') == ''))
comprobar('alumnado con dos especialidades = 1', control['Alumnado con dos especialidades'] == 1,
          control['Alumnado con dos especialidades'])
# Solo quien lleva la bandurria: la fila de curso «1» no tiene etapa y no entra.
comprobar('alumnado de profesional sin especialidad = 1',
          control['Alumnado de profesional sin especialidad reconocida'] == 1,
          control['Alumnado de profesional sin especialidad reconocida'])
comprobar('una pendiente, que son dos filas (OR y EX)', control['Pendientes'] == 2, control['Pendientes'])
comprobar('la asignatura que no está en el catálogo cuenta una vez, no dos',
          control['Asignaturas que no están en el catálogo de su etapa'] == 1,
          control['Asignaturas que no están en el catálogo de su etapa'])

inc = libro['INCIDENCIAS']
print('\nINCIDENCIAS:')
for fila_ in inc.iter_rows(min_row=3, max_row=12, values_only=True):
    if any(x not in (None, '') for x in fila_):
        print('   ', [x for x in fila_ if x not in (None, '')])
texto = ' '.join(str(x) for fila_ in inc.iter_rows(values_only=True) for x in fila_ if x is not None)
comprobar('INCIDENCIAS lista la bandurria', 'BANDURRIA' in texto)
comprobar('INCIDENCIAS lista el curso «1»', ' 1 ' in ' %s ' % texto)
comprobar('INCIDENCIAS lista la nota NP', 'NP' in texto)
comprobar('INCIDENCIAS lista la evaluación 3EV', '3EV' in texto)
comprobar('INCIDENCIAS no lleva ningún NIA ni nombre',
          not any(('NOMBRE%s' % n) in texto or ('APELLIDO%s' % n) in texto for n in ids)
          and not any(str(n) in texto.split() for n in ids))

print('\n%s' % ('Todo cuadra.' if not fallos else 'FALLAN %d.' % fallos))
sys.exit(1 if fallos else 0)
